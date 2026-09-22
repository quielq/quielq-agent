"""create_docx: generate a .docx file into the agent's own memory folder.

Needs no credentials, unlike the Google Docs path (still to come, once the
shared Google OAuth setup `daily-brief` also needs exists - see the plan
doc). `drafts/` is a purpose-specific subfolder alongside the standard
profile/knowledge/procedures/experiences taxonomy, same pattern as
`pending_actions/` in tools/approval.py.
"""

from __future__ import annotations

from docx import Document

from quielq_agent.tools import ToolContext

CREATE_DOCX_SCHEMA = {
    "type": "function",
    "function": {
        "name": "create_docx",
        "description": (
            "Create a .docx file with the given title and body text, saved to "
            "this agent's own memory folder. Paragraphs in the body should be "
            "separated by blank lines."
        ),
        "parameters": {
            "type": "object",
            "properties": {
                "filename": {"type": "string", "description": "Filename without extension, e.g. 'my-draft'."},
                "title": {"type": "string", "description": "Document title, used as a heading."},
                "body": {"type": "string", "description": "The document body text."},
            },
            "required": ["filename", "title", "body"],
        },
    },
}


def create_docx(filename: str, title: str, body: str, context: ToolContext | None = None) -> str:
    if context is None or context.memory_dir is None:
        return "error: create_docx needs an agent with a memory_dir configured"

    safe_name = "".join(c for c in filename if c.isalnum() or c in "-_") or "document"
    drafts_dir = context.memory_dir / "drafts"
    drafts_dir.mkdir(parents=True, exist_ok=True)
    path = drafts_dir / f"{safe_name}.docx"

    document = Document()
    document.add_heading(title, level=1)
    for paragraph in body.split("\n\n"):
        if paragraph.strip():
            document.add_paragraph(paragraph.strip())
    document.save(path)

    return f"Saved to {path.relative_to(context.memory_dir)} ({len(body)} chars in body)."
